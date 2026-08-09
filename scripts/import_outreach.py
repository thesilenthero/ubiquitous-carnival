#!/usr/bin/env python3
"""
Import the "Job Search Tracker" Outreach tab into the app's contacts tables.

Mapping (Outreach sheet):
  Name                 -> name
  Company              -> company
  Type                 -> relationship
  Status               -> status (Pending / Connected / Followed up / Closed)
  Date Contacted       -> 'outreach' interaction (if date present)
  Most recent response -> 'response' interaction (if date present)
  Connect Date         -> 'connected' interaction (if date present)
  Follow-up by         -> next_action_date (next_action defaults to "Follow up")
  Next steps           -> next_action (overrides the default)
  Notes                -> notes

Destructive for contacts ONLY: it clears contacts + interactions and reloads
them from the sheet. Applications and stage history are never touched, so it is
safe to re-run without losing job_url/source/etc. hand edits.

Usage:
  python3 scripts/import_outreach.py "/path/to/Job Search Tracker.xlsx" [data/app.db]
"""
import os
import sys
import sqlite3
import secrets
import datetime as dt

import openpyxl

ALPHABET = "useandom-26T198340PX75pxJACKVERYMINDBUSHWOLFGQZbfghjklqvwyzrict"


def nano(n: int = 21) -> str:
    return "".join(secrets.choice(ALPHABET) for _ in range(n))


def iso_at(d: dt.datetime) -> str:
    """Midnight-UTC ISO timestamp for a sheet date (matches app convention)."""
    return dt.datetime(d.year, d.month, d.day, tzinfo=dt.timezone.utc).isoformat().replace("+00:00", "Z")


def date_str(d: dt.datetime) -> str:
    return f"{d.year:04d}-{d.month:02d}-{d.day:02d}"


def main() -> None:
    here = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    xlsx = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
        os.path.expanduser("~"), "Downloads", "Job Search Tracker-5.xlsx"
    )
    db_path = sys.argv[2] if len(sys.argv) > 2 else os.path.join(here, "data", "app.db")

    if not os.path.exists(xlsx):
        sys.exit(f"Spreadsheet not found: {xlsx}")
    if not os.path.exists(db_path):
        sys.exit(f"Database not found: {db_path} (start the app once to create it)")

    wb = openpyxl.load_workbook(xlsx, data_only=True)
    if "Outreach" not in wb.sheetnames:
        sys.exit("No Outreach sheet in this workbook — nothing to import.")
    ws = wb["Outreach"]
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r and r[0] not in (None, "")]

    now_iso = dt.datetime.now(dt.timezone.utc).isoformat().replace("+00:00", "Z")

    conn = sqlite3.connect(db_path)
    conn.execute("PRAGMA foreign_keys = ON")
    cur = conn.cursor()
    cur.execute("DELETE FROM interactions")
    cur.execute("DELETE FROM contacts")

    n_contacts = 0
    n_interactions = 0

    for r in rows:
        (name, company, rel_type, status, contacted,
         response, follow_up_by, connect_date, next_steps, notes) = (list(r) + [None] * 10)[:10]

        contact_id = nano()

        interactions = []
        if isinstance(contacted, dt.datetime):
            interactions.append(("outreach", iso_at(contacted)))
        if isinstance(connect_date, dt.datetime):
            interactions.append(("connected", iso_at(connect_date)))
        if isinstance(response, dt.datetime):
            interactions.append(("response", iso_at(response)))

        next_action = str(next_steps).strip() if next_steps else (
            "Follow up" if isinstance(follow_up_by, dt.datetime) else None
        )
        next_action_date = date_str(follow_up_by) if isinstance(follow_up_by, dt.datetime) else None

        created = interactions[0][1] if interactions else now_iso
        updated = interactions[-1][1] if interactions else now_iso

        cur.execute(
            """INSERT INTO contacts
               (id, name, company, role_title, relationship, status, email,
                linkedin_url, next_action, next_action_date, notes,
                created_at, updated_at)
               VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (contact_id, str(name).strip(), str(company).strip() if company else None,
             None, str(rel_type).strip() if rel_type else None,
             str(status).strip() if status else "Pending", None, None,
             next_action, next_action_date,
             str(notes).strip() if notes else None, created, updated),
        )
        for kind, when in interactions:
            cur.execute(
                "INSERT INTO interactions (id, contact_id, application_id, kind, note, occurred_at) VALUES (?,?,?,?,?,?)",
                (nano(), contact_id, None, kind, None, when),
            )
            n_interactions += 1
        n_contacts += 1

    conn.commit()
    conn.close()

    print(f"Imported {n_contacts} contacts, {n_interactions} interactions.")


if __name__ == "__main__":
    main()
